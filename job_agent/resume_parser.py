from io import BytesIO
from pathlib import Path


def extract_resume_text(file_bytes: bytes, filename: str) -> str:
    """Extract text from PDF or DOCX uploaded through Streamlit."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        import fitz  # PyMuPDF
        text_parts = []
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            for page in doc:
                text_parts.append(page.get_text("text"))
        return "\n".join(text_parts).strip()
    if suffix == ".docx":
        from docx import Document
        doc = Document(BytesIO(file_bytes))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()
    raise ValueError("Unsupported file type. Please upload a PDF or DOCX file.")
